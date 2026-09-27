"""Excel report generation (openpyxl) for supervisor attendance exports."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)


def build_overview_workbook(rollups, from_date: date, to_date: date, dept_name=None) -> BytesIO:
    """Build an .xlsx of per-employee attendance rollups. Returns a BytesIO stream."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    scope = dept_name or "All departments"
    ws["A1"] = f"Attendance report — {scope}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Period: {from_date.isoformat()} to {to_date.isoformat()}"

    headers = ["Employee Code", "Name", "Department", "Present", "Absent", "Attendance %", "Total Hours"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    r = header_row + 1
    for item in rollups:
        ws.cell(row=r, column=1, value=item.external_id)
        ws.cell(row=r, column=2, value=item.display_name)
        ws.cell(row=r, column=3, value=item.department_name or "—")
        ws.cell(row=r, column=4, value=item.present)
        ws.cell(row=r, column=5, value=item.absent)
        ws.cell(row=r, column=6, value=item.attendance_pct)
        ws.cell(row=r, column=7, value=item.total_hours)
        r += 1

    # Column widths for readability.
    widths = [16, 22, 16, 10, 9, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
